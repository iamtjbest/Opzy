"""Load opportunities from the source tracking sheet.

The sheet (docs/OPZY_SOURCE_TRACKING.xlsx) is the source of truth for hand-entered
opportunities; its columns map 1:1 to the `opportunities` table.

Run from the backend/ directory:

    python -m scripts.seed_opportunities          # insert new rows into DATABASE_URL
    python -m scripts.seed_opportunities --sql    # regenerate docs/seed_opportunities_dev.sql

Both are safe to re-run: a row whose title and organization are already in the table is
updated to match the sheet (and left alone if nothing changed). Inserting refuses a
non-local DATABASE_URL unless --allow-nonlocal is passed, because the sheet currently holds
fictional dev data.
"""

import argparse
import asyncio
import sys
from datetime import UTC, date, datetime
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.countries import is_country
from app.models import Opportunity
from app.models.base import EDUCATION_LEVELS, OPPORTUNITY_STATUSES, OPPORTUNITY_TYPES

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SHEET = REPO_ROOT / "docs" / "OPZY_SOURCE_TRACKING.xlsx"
DEFAULT_SQL = REPO_ROOT / "docs" / "seed_opportunities_dev.sql"
SHEET_NAME = "Opportunities"
NOT_CONFIRMED = "not confirmed"
LOCAL_HOSTS = {"localhost", "127.0.0.1", "::1", "db"}

# Sheet header -> table column. "Notes" is for people reading the sheet and isn't imported.
COLUMNS = {
    "Title": "title",
    "Organization": "organization",
    "Category": "category",
    "Geography": "geography",
    "Description": "description",
    "Deadline": "deadline",
    "Eligibility Notes": "eligibility_notes",
    "Application URL": "application_url",
    "Source URL": "source_url",
    "Quality Rating": "quality_rating",
    "Verified?": "verified",
    "Status": "status",
    "Eligible Countries": "eligible_countries",
    "Education Levels": "education_levels",
    "Fields of Study": "fields_of_study",
    "Skills": "skills",
}
HEADERS = [*COLUMNS, "Notes"]

# Postgres types for the non-text columns, so NULLs in generated SQL are typed.
SQL_TYPES = {"deadline": "date", "quality_rating": "integer", "verified": "boolean"}

# Shorthands allowed in Eligible Countries. ECOWAS membership as of 2025, after Burkina
# Faso, Mali and Niger left in January 2025.
REGIONS = {
    "ECOWAS": ("BJ", "CI", "CV", "GH", "GM", "GN", "GW", "LR", "NG", "SL", "SN", "TG"),
}


def _text(value: object) -> str | None:
    if value is None:
        return None
    return str(value).strip() or None


def _deadline(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if text is None or text.lower() == NOT_CONFIRMED:
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        raise ValueError(
            f"Deadline must be a date or NOT CONFIRMED, got {text!r}"
        ) from None


def _rating(value: object) -> int | None:
    if value is None or _text(value) is None:
        return None
    if isinstance(value, bool) or not isinstance(value, int | float) or value != int(value):
        raise ValueError(f"Quality Rating must be a whole number 1-5, got {value!r}")
    if not 1 <= value <= 5:
        raise ValueError(f"Quality Rating must be 1-5, got {value!r}")
    return int(value)


def _verified(value: object) -> bool:
    text = (_text(value) or "no").lower()
    if text not in ("yes", "no"):
        raise ValueError(f"Verified? must be Yes or No, got {value!r}")
    return text == "yes"


def _choice(value: object, header: str, allowed: tuple[str, ...], default: str | None) -> str:
    text = (_text(value) or default or "").lower()
    if text not in allowed:
        raise ValueError(f"{header} must be one of {', '.join(allowed)}, got {value!r}")
    return text


def _items(value: object) -> list[str]:
    """A comma-separated cell as a list: trimmed, blanks dropped, deduped ignoring case."""
    seen: dict[str, str] = {}
    for part in (_text(value) or "").split(","):
        part = part.strip()
        if part:
            seen.setdefault(part.lower(), part)
    return list(seen.values())


def _countries(value: object) -> list[str]:
    codes: set[str] = set()
    for item in _items(value):
        expanded = REGIONS.get(item.upper(), (item.upper(),))
        if not all(is_country(code) for code in expanded):
            raise ValueError(
                f"Eligible Countries must be ISO country codes (NG, GH, …) or "
                f"{', '.join(REGIONS)}, got {item!r}"
            )
        codes.update(expanded)
    return sorted(codes)


def _levels(value: object) -> list[str]:
    levels = [item.lower() for item in _items(value)]
    for level in levels:
        if level not in EDUCATION_LEVELS:
            raise ValueError(
                f"Education Levels must be from {', '.join(EDUCATION_LEVELS)}, got {level!r}"
            )
    return sorted(levels, key=EDUCATION_LEVELS.index)


def _parse_row(cells: dict[str, object]) -> dict:
    title = _text(cells["Title"])
    if title is None:
        raise ValueError("Title is required")
    return {
        "title": title,
        "organization": _text(cells["Organization"]),
        "category": _choice(cells["Category"], "Category", OPPORTUNITY_TYPES, None),
        "geography": _text(cells["Geography"]),
        "description": _text(cells["Description"]),
        "deadline": _deadline(cells["Deadline"]),
        "eligibility_notes": _text(cells["Eligibility Notes"]),
        "application_url": _text(cells["Application URL"]),
        "source_url": _text(cells["Source URL"]),
        "quality_rating": _rating(cells["Quality Rating"]),
        "verified": _verified(cells["Verified?"]),
        "status": _choice(cells["Status"], "Status", OPPORTUNITY_STATUSES, "active"),
        "eligible_countries": _countries(cells["Eligible Countries"]),
        "education_levels": _levels(cells["Education Levels"]),
        "fields_of_study": _items(cells["Fields of Study"]),
        "skills": _items(cells["Skills"]),
    }


def parse_sheet(path: Path) -> list[dict]:
    """The sheet's rows as `opportunities` column dicts. Raises ValueError naming the row."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        rows = ws.iter_rows(values_only=True)
        header = [_text(h) for h in next(rows, ())]
        missing = [h for h in COLUMNS if h not in header]
        if missing:
            raise ValueError(f"Sheet is missing column(s): {', '.join(missing)}")
        index = {h: header.index(h) for h in COLUMNS}

        parsed = []
        for number, values in enumerate(rows, start=2):
            if all(_text(v) is None for v in values):
                continue
            cells = {h: values[i] if i < len(values) else None for h, i in index.items()}
            try:
                parsed.append(_parse_row(cells))
            except ValueError as exc:
                raise ValueError(f"Row {number}: {exc}") from None
        return parsed
    finally:
        wb.close()


async def seed(db: AsyncSession, rows: list[dict]) -> tuple[int, int]:
    """Insert new rows and update changed ones, matched by title + organization.

    Returns (added, updated). Rows already matching the sheet are left alone.
    """
    existing = {(o.title, o.organization): o for o in await db.scalars(select(Opportunity))}
    added = updated = 0
    for row in rows:
        key = (row["title"], row["organization"])
        current = existing.get(key)
        if current is None:
            existing[key] = Opportunity(**row)
            db.add(existing[key])
            added += 1
        elif any(getattr(current, column) != value for column, value in row.items()):
            for column, value in row.items():
                setattr(current, column, value)
            current.updated_at = datetime.now(UTC)
            updated += 1
    await db.flush()
    return added, updated


def _literal(column: str, value: object) -> str:
    if value is None:
        sql_type = SQL_TYPES.get(column)
        return f"null::{sql_type}" if sql_type else "null"
    if isinstance(value, list):
        if not value:
            return "'{}'::text[]"
        return "array[" + ", ".join(_literal(column, v) for v in value) + "]::text[]"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, date):
        return f"date '{value.isoformat()}'"
    return "'" + str(value).replace("'", "''") + "'"


def to_sql(rows: list[dict]) -> str:
    """Idempotent upserts for psql / the Supabase SQL editor. Run after `alembic upgrade head`."""
    columns = list(COLUMNS.values())
    parts = [
        f"-- GENERATED from docs/{DEFAULT_SHEET.name} by backend/scripts/seed_opportunities.py.\n"
        "-- Don't edit by hand: edit the sheet, then run `python -m scripts.seed_opportunities --sql`.\n"
        "--\n"
        "-- DEV/TEST DATA: includes fictional opportunities (see the sheet's Notes column).\n"
        "-- Never run this against production.\n"
        "--\n"
        "-- Needs the migrated schema: run after `alembic upgrade head`.\n"
        "-- Safe to re-run: rows already present (same title and organization) are updated to\n"
        "-- match the sheet; unchanged rows are left alone.\n"
    ]
    for row in rows:
        literals = [_literal(c, row[c]) for c in columns]
        title = _literal("title", row["title"])
        organization = _literal("organization", row["organization"])
        assignments = ",\n    ".join(f"{c} = {v}" for c, v in zip(columns, literals))
        values = ",\n       ".join(literals)
        parts.append(
            # Update first, and only when something differs, so a re-run doesn't bump
            # updated_at on unchanged rows.
            f"\nupdate opportunities set\n"
            f"    {assignments},\n"
            f"    updated_at = now()\n"
            f"where title = {title}\n"
            f"  and organization is not distinct from {organization}\n"
            f"  and ({', '.join(columns)}) is distinct from ({', '.join(literals)});\n"
            f"\ninsert into opportunities ({', '.join(columns)})\n"
            f"select {values}\n"
            f"where not exists (\n"
            f"  select 1 from opportunities\n"
            f"  where title = {title}\n"
            f"    and organization is not distinct from {organization}\n"
            f");\n"
        )
    return "".join(parts)


async def _insert(rows: list[dict]) -> tuple[int, int]:
    from app.core.db import SessionLocal, engine

    try:
        async with SessionLocal() as session:
            result = await seed(session, rows)
            await session.commit()
            return result
    finally:
        await engine.dispose()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--sheet", type=Path, default=DEFAULT_SHEET)
    parser.add_argument(
        "--sql",
        nargs="?",
        type=Path,
        const=DEFAULT_SQL,
        help=f"write SQL instead of inserting (default path: {DEFAULT_SQL.relative_to(REPO_ROOT)})",
    )
    parser.add_argument(
        "--allow-nonlocal", action="store_true", help="insert into a non-local database"
    )
    args = parser.parse_args()

    try:
        rows = parse_sheet(args.sheet)
    except ValueError as exc:
        sys.exit(f"{args.sheet}: {exc}")

    if args.sql:
        args.sql.write_text(to_sql(rows))
        print(f"Wrote {len(rows)} opportunities to {args.sql}")
        return

    from app.core.config import get_settings

    host = urlsplit(get_settings().database_url).hostname
    if host not in LOCAL_HOSTS and not args.allow_nonlocal:
        sys.exit(
            f"Refusing to seed non-local database host {host!r}: the sheet holds fictional "
            "dev data. Pass --allow-nonlocal if you really mean it."
        )
    added, updated = asyncio.run(_insert(rows))
    print(f"Inserted {added} and updated {updated} of {len(rows)} opportunities")


if __name__ == "__main__":
    main()
